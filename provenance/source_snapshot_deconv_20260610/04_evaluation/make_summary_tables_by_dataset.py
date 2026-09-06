#!/usr/bin/env python
"""Per-dataset method-comparison tables, built ONLY from summary_all_methods.csv.

One table per dataset (3 datasets -> 3 CSVs + 3 sheets in one xlsx). Each table is a
per-METHOD summary (summary_all_methods.csv is aggregated across cell types, so there is
no per-cell-type breakdown to put here). Layout, per the request:

  [ left block ]                                                 (blank cols)   [ right block ]
  Method | profile-wise (all/marker) | profile-wise sxct (all/marker)           Method | gene-wise (all/marker)
         | sample-wise  (all/marker) | pseudo-bulk   (all/marker)

For EVERY metric we report BOTH the mean and the median (summary_all_methods.csv carries
both columns), so the boxplot figures — whose central line is the median while the annotated
number / diamond is the mean — can be cross-checked against the table either way. Column
names are suffixed " [mean]" / " [median]" accordingly.

Values = the `mean` / `median` columns of summary_all_methods.csv (aggregated over the axis's
units: cell types for profile/pseudobulk, samples for sample, sample x cell type for
sample_profile, genes for gene-wise). Blank where that axis is absent / n=0 for a method
(e.g. sample & sample_profile on GSE159585 which has no sample dimension).
"""
import os
import numpy as np, pandas as pd

ROOT = os.environ.get("WORK_ROOT", "/disk1/maijl/deconv/deconv_20260605") + "/03_result_tables"
SUMMARY = f"{ROOT}/summary_all_methods.csv"
METHODS = ["DeconvSC", "BayesPrism", "CIBERSORTx", "DISSECT", "TAPE"]   # canonical labels (cVAE -> DeconvSC)
DATASETS = ["HCA_fold2", "GSE141115", "GSE159585"]
ND = 4  # round


def _pair(base):
    """('<base> [mean]', '<base> [median]') header pair (descriptor stays in (), stat in [])."""
    return (f"{base} [mean]", f"{base} [median]")


# base metric names (all-gene, marker-gene)
B_PROF_CT   = ("All gene profile-wise PCC (per cell type)",        "Marker gene profile-wise PCC (per cell type)")
B_PROF_SXCT = ("All gene profile-wise PCC (per sample×cell type)", "Marker gene profile-wise PCC (per sample×cell type)")
B_SAMPLE    = ("All gene sample-wise PCC",                         "Marker gene sample-wise PCC")
B_PB        = ("All gene pseudo-bulk PCC",                         "Marker gene pseudo-bulk PCC")
B_GENE      = ("All gene gene-wise PCC (common set)",             "Marker gene gene-wise PCC")  # all-gene = cross-method common set (values from summary_all_methods, fixed by apply_common_genewise.py)


def main():
    S = pd.read_csv(SUMMARY)
    # accept the legacy 'cVAE' method label and emit the canonical 'DeconvSC' (scheme-agnostic:
    # works whether summary_all_methods.csv already says DeconvSC or still says cVAE)
    S["method"] = S["method"].replace({"cVAE": "DeconvSC"})

    def lk(ds, m, gset, axis):
        """(n, mean, median) for one cell of summary; (0, nan, nan) if absent/empty."""
        r = S[(S.dataset == ds) & (S.method == m) & (S.gene_set == gset) & (S.axis == axis)]
        if r.empty:
            return 0, np.nan, np.nan
        n = int(r["n"].iloc[0]); mean = r["mean"].iloc[0]; med = r["median"].iloc[0]
        if n <= 0 or pd.isna(mean):
            return 0, np.nan, np.nan
        return n, round(float(mean), ND), round(float(med), ND)

    def val(ds, m, gset, axis):
        """(mean, median) for a directly-named axis."""
        _, mean, med = lk(ds, m, gset, axis)
        return mean, med

    def profiles(ds, m, gset):
        """((per_ct_mean, per_ct_median), (per_sxct_mean, per_sxct_median)) — routes the
        ambiguous `profile` axis by comparing its n to the pseudobulk n (which is per cell type),
        carrying mean+median together. GSE141115 competitors store per-(sample×celltype) values
        under axis name `profile` (n >> pseudobulk n); cVAE stores per-cell-type under `profile`
        and per-(sample×celltype) under `sample_profile`. HCA has both axes per method."""
        pb_n, _, _ = lk(ds, m, gset, "pseudobulk")
        pr_n, pr_mean, pr_med = lk(ds, m, gset, "profile")
        sp_n, sp_mean, sp_med = lk(ds, m, gset, "sample_profile")
        per_ct, per_sxct = (np.nan, np.nan), (np.nan, np.nan)
        if pr_n > 0:
            if pb_n > 0 and pr_n > pb_n * 1.5:
                per_sxct = (pr_mean, pr_med)      # `profile` axis is actually per (sample×cell type)
            else:
                per_ct = (pr_mean, pr_med)        # `profile` axis is per cell type
        if sp_n > 0:
            per_sxct = (sp_mean, sp_med)          # canonical per (sample×cell type)
        return per_ct, per_sxct

    out_tables = {}
    for ds in DATASETS:
        left_cols = (["Method"]
                     + list(_pair(B_PROF_CT[0])) + list(_pair(B_PROF_CT[1]))
                     + list(_pair(B_PROF_SXCT[0])) + list(_pair(B_PROF_SXCT[1]))
                     + list(_pair(B_SAMPLE[0])) + list(_pair(B_SAMPLE[1]))
                     + list(_pair(B_PB[0])) + list(_pair(B_PB[1])))
        left = []
        for m in METHODS:
            pct_a, psx_a = profiles(ds, m, "all_genes")
            pct_m, psx_m = profiles(ds, m, "marker_genes")
            sa, sm = val(ds, m, "all_genes", "sample"),     val(ds, m, "marker_genes", "sample")
            pba, pbm = val(ds, m, "all_genes", "pseudobulk"), val(ds, m, "marker_genes", "pseudobulk")
            left.append([m, *pct_a, *pct_m, *psx_a, *psx_m, *sa, *sm, *pba, *pbm])
        left_df = pd.DataFrame(left, columns=left_cols)

        # right block (gene-wise): mean + median
        right_cols = ["Method"] + list(_pair(B_GENE[0])) + list(_pair(B_GENE[1]))
        right = []
        for m in METHODS:
            ga, gm = val(ds, m, "all_genes", "gene"), val(ds, m, "marker_genes", "gene")
            right.append([m, *ga, *gm])
        right_df = pd.DataFrame(right, columns=right_cols)

        # side-by-side with two blank spacer columns
        blank = pd.DataFrame({" ": [""] * len(METHODS), "  ": [""] * len(METHODS)})
        out = pd.concat([left_df, blank, right_df], axis=1)
        out_tables[ds] = out
        out.to_csv(f"{ROOT}/summary_table_{ds}.csv", index=False)
        print(f"  {ds:12s} -> summary_table_{ds}.csv  ({out.shape[0]} methods x {out.shape[1]} cols)")

    with pd.ExcelWriter(f"{ROOT}/summary_tables_by_dataset.xlsx", engine="openpyxl") as xw:
        for ds in DATASETS:
            out_tables[ds].to_excel(xw, sheet_name=ds[:31], index=False)
        # widen columns + clear the spacer/duplicate headers for readability
        from openpyxl.utils import get_column_letter
        for ds in DATASETS:
            wsv = xw.sheets[ds[:31]]
            for ci, col in enumerate(out_tables[ds].columns, start=1):
                wsv.column_dimensions[get_column_letter(ci)].width = 15 if ci > 1 else 12
                # blank out spacer headers (" ", "  ")
                if str(col).strip() == "":
                    wsv.cell(row=1, column=ci).value = None
    print(f"[done] -> {ROOT}/summary_tables_by_dataset.xlsx (+ 3 summary_table_<ds>.csv)")


if __name__ == "__main__":
    main()
